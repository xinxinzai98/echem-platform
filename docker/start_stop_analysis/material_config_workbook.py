#!/usr/bin/env python3
"""OpenPyXL-compatible material configuration workbook adapter.

This is a public-dependency replacement for ``material_config_workbook.mjs``.
It intentionally keeps the same three runtime CLI modes and JSON contracts:

    update <snapshot.json> <workbook.xlsx>
    read <workbook.xlsx> <output.json>
    apply-json <snapshot.json> <config.json> <workbook.xlsx>

The analysis pipeline treats dataset and material fingerprints as opaque values;
this module only preserves and validates them.
"""

from __future__ import annotations

import json
import math
import os
import sys
import tempfile
from pathlib import Path, PurePosixPath
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo


COLORS = {
    "navy": "17365D",
    "blue": "2F75B5",
    "pale_blue": "DDEBF7",
    "green": "548235",
    "pale_green": "E2F0D9",
    "orange": "C65911",
    "pale_orange": "FCE4D6",
    "red": "C00000",
    "pale_red": "F4CCCC",
    "yellow": "FFF2CC",
    "gray": "666666",
    "pale_gray": "E7E6E6",
    "border": "B4C6E7",
    "white": "FFFFFF",
    "text": "1F1F1F",
    "body_border": "D9E2F3",
}

MATERIAL_HEADERS = [
    "材料键（勿改）",
    "自动识别名称",
    "绘图名称（可修改）",
    "进入总结图集（是/否）",
    "配置检查",
    "数据状态",
    "启停文件数",
    "总数据点",
    "接续顺序",
    "材料相对路径",
    "数据指纹（勿改）",
    "用户备注（可填写）",
    "最近扫描时间",
    "测试类型",
    "阴极电流中位数 / A cm⁻²",
    "恢复段电流中位数 / A cm⁻²",
]

FILE_HEADERS = [
    "文件相对路径",
    "材料键",
    "文件名",
    "文件名类型",
    "是否纳入启停",
    "状态或排除原因",
    "有效数据点",
    "电流档位 / A cm⁻²",
    "起始时间 / s",
    "结束时间 / s",
    "文件修改时间",
    "SHA-256",
    "测试类型",
    "阴极电流中位数 / A cm⁻²",
    "恢复段电流中位数 / A cm⁻²",
]

THIN_SIDE = Side(style="thin", color=COLORS["border"])
BODY_SIDE = Side(style="thin", color=COLORS["body_border"])
THIN_BORDER = Border(
    left=THIN_SIDE,
    right=THIN_SIDE,
    top=THIN_SIDE,
    bottom=THIN_SIDE,
)
BODY_BORDER = Border(
    left=BODY_SIDE,
    right=BODY_SIDE,
    top=BODY_SIDE,
    bottom=BODY_SIDE,
)


def normalize_cell(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def safe_number(value: Any) -> int | float | None:
    """Return a real finite number, keeping missing values visibly blank."""
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, str) and not value.strip():
        return None
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    if not math.isfinite(number):
        return None
    if isinstance(value, int):
        return value
    return number


def bool_zh(value: Any) -> str:
    return "是" if bool(value) else "否"


def set_text(cell, value: Any) -> None:
    """Write literal text even when it starts with an Excel formula marker."""
    if value is None:
        cell.value = None
        return
    cell.value = str(value)
    cell.data_type = "s"


def fill(color: str) -> PatternFill:
    return PatternFill(fill_type="solid", fgColor=color)


def apply_cells(worksheet, cell_range: str, callback) -> None:
    for row in worksheet[cell_range]:
        for cell in row:
            callback(cell)


def style_grid_cell(cell, *, body: bool = False) -> None:
    cell.border = BODY_BORDER if body else THIN_BORDER
    cell.alignment = Alignment(vertical="center", wrap_text=True)


def set_column_widths(worksheet, widths: dict[str, float]) -> None:
    for column, width in widths.items():
        worksheet.column_dimensions[column].width = width


def load_json(path: Path) -> dict[str, Any]:
    payload = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(payload, dict):
        raise ValueError(f"JSON 顶层必须是对象: {path}")
    return payload


def atomic_write_json(path: Path, payload: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, temporary_name = tempfile.mkstemp(
        prefix=f".{path.name}.",
        suffix=".tmp",
        dir=path.parent,
    )
    temporary_path = Path(temporary_name)
    try:
        with os.fdopen(fd, "w", encoding="utf-8") as handle:
            json.dump(payload, handle, ensure_ascii=False, indent=2)
            handle.write("\n")
            handle.flush()
            os.fsync(handle.fileno())
        os.replace(temporary_path, path)
    finally:
        temporary_path.unlink(missing_ok=True)


def atomic_save_workbook(workbook: Workbook, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, temporary_name = tempfile.mkstemp(
        prefix=f".{path.name}.",
        suffix=".tmp.xlsx",
        dir=path.parent,
    )
    os.close(fd)
    temporary_path = Path(temporary_name)
    try:
        workbook.save(temporary_path)
        os.replace(temporary_path, path)
    finally:
        temporary_path.unlink(missing_ok=True)


def find_material_header_row(worksheet) -> int | None:
    for row_index in range(1, worksheet.max_row + 1):
        if normalize_cell(worksheet.cell(row_index, 1).value) == MATERIAL_HEADERS[0]:
            return row_index
    return None


def load_existing_config(workbook_path: Path) -> dict[str, Any]:
    if not workbook_path.exists():
        return {"rows_by_key": {}, "dataset_fingerprint": ""}

    workbook = load_workbook(workbook_path, data_only=False, read_only=False)
    rows_by_key: dict[str, dict[str, str]] = {}
    dataset_fingerprint = ""
    try:
        if "使用说明" in workbook.sheetnames:
            dataset_fingerprint = normalize_cell(workbook["使用说明"]["B20"].value)
        if "材料配置" in workbook.sheetnames:
            worksheet = workbook["材料配置"]
            header_row = find_material_header_row(worksheet)
            if header_row is not None:
                for row_index in range(header_row + 1, worksheet.max_row + 1):
                    key = normalize_cell(worksheet.cell(row_index, 1).value)
                    if not key:
                        continue
                    rows_by_key[key] = {
                        "key": key,
                        "auto_name": normalize_cell(worksheet.cell(row_index, 2).value),
                        "plot_name": normalize_cell(worksheet.cell(row_index, 3).value),
                        "include": normalize_cell(worksheet.cell(row_index, 4).value),
                        "status": normalize_cell(worksheet.cell(row_index, 6).value),
                        "fingerprint": normalize_cell(worksheet.cell(row_index, 11).value),
                        "notes": normalize_cell(worksheet.cell(row_index, 12).value),
                    }
    finally:
        workbook.close()
    return {
        "rows_by_key": rows_by_key,
        "dataset_fingerprint": dataset_fingerprint,
    }


def posix_basename(value: str) -> str:
    name = PurePosixPath(value).name
    return name or value


def build_material_rows(
    snapshot: dict[str, Any],
    existing: dict[str, Any],
) -> list[dict[str, Any]]:
    materials = snapshot.get("materials")
    if not isinstance(materials, list):
        raise ValueError("材料快照缺少 materials 数组。")
    now = snapshot.get("generated_at", "")
    existing_by_key = existing.get("rows_by_key", {})
    current_keys = {
        normalize_cell(item.get("key"))
        for item in materials
        if isinstance(item, dict)
    }
    rows: list[dict[str, Any]] = []

    for material in materials:
        if not isinstance(material, dict):
            raise ValueError("材料快照中的材料记录必须是对象。")
        key = normalize_cell(material.get("key"))
        if not key:
            raise ValueError("材料快照包含空材料键。")
        previous = existing_by_key.get(key)
        previous_include = normalize_cell(previous.get("include")) if previous else ""
        include = previous_include if previous_include in {"是", "否"} else "是"
        plot_name = (
            normalize_cell(previous.get("plot_name")) if previous else ""
        ) or normalize_cell(material.get("auto_name"))
        fingerprint = normalize_cell(material.get("fingerprint"))
        if previous is None:
            status = "新增"
        elif normalize_cell(previous.get("fingerprint")) == fingerprint:
            status = "未变化"
        else:
            status = "数据已更新"
        source_files = material.get("ordered_source_files")
        if not isinstance(source_files, list):
            source_files = []
        rows.append(
            {
                "key": key,
                "auto_name": normalize_cell(material.get("auto_name")),
                "plot_name": plot_name,
                "include": include,
                "status": status,
                "standard_file_count": safe_number(material.get("standard_file_count")),
                "total_data_points": safe_number(material.get("total_data_points")),
                "order": " → ".join(normalize_cell(item) for item in source_files),
                "relative_path": key,
                "fingerprint": fingerprint,
                "notes": normalize_cell(previous.get("notes")) if previous else "",
                "scanned_at": normalize_cell(now),
                "test_types": normalize_cell(material.get("test_types")),
                "cathodic_current": safe_number(
                    material.get("cathodic_current_median_a_cm2")
                ),
                "recovery_current": safe_number(
                    material.get("recovery_current_median_a_cm2")
                ),
            }
        )

    for key, previous in existing_by_key.items():
        if key in current_keys:
            continue
        auto_name = normalize_cell(previous.get("auto_name")) or posix_basename(key)
        rows.append(
            {
                "key": key,
                "auto_name": auto_name,
                "plot_name": normalize_cell(previous.get("plot_name")) or auto_name,
                "include": "否",
                "status": "源文件已移除",
                "standard_file_count": 0,
                "total_data_points": 0,
                "order": "",
                "relative_path": key,
                "fingerprint": normalize_cell(previous.get("fingerprint")),
                "notes": normalize_cell(previous.get("notes")),
                "scanned_at": normalize_cell(now),
                "test_types": "",
                "cathodic_current": None,
                "recovery_current": None,
            }
        )

    return sorted(rows, key=lambda item: item["key"])


def apply_title_band(worksheet, cell_range: str, title: str) -> None:
    worksheet.merge_cells(cell_range)
    cell = worksheet[cell_range.split(":", 1)[0]]
    set_text(cell, title)
    cell.fill = fill(COLORS["navy"])
    cell.font = Font(bold=True, color=COLORS["white"], size=18)
    cell.alignment = Alignment(horizontal="left", vertical="center")
    worksheet.row_dimensions[cell.row].height = 34


def build_usage_sheet(
    workbook: Workbook,
    snapshot: dict[str, Any],
    row_count: int,
) -> None:
    worksheet = workbook.create_sheet("使用说明")
    worksheet.sheet_view.showGridLines = False
    apply_title_band(worksheet, "A1:H1", "启停绘图材料配置")

    worksheet.merge_cells("A3:H3")
    set_text(
        worksheet["A3"],
        "先更新文件和配置表，再由你确认名称与是否入图；确认前不会生成新的总结图集。",
    )
    worksheet["A3"].fill = fill(COLORS["pale_blue"])
    worksheet["A3"].font = Font(color=COLORS["navy"], bold=True, size=11)
    worksheet["A3"].alignment = Alignment(wrap_text=True, vertical="center")
    worksheet.row_dimensions[3].height = 34

    labels = ["当前材料", "选入图集", "不进入图集", "需要处理"]
    steps = [
        ("步骤 1", "双击“1_更新材料配置表.command”"),
        ("步骤 2", "在“材料配置”中修改黄色列并保存"),
        ("步骤 3", "双击“2_按配置生成图集.command”"),
        ("安全门", "若更新表后源文件又变化，绘图会停止并要求重新执行步骤 1"),
    ]
    for row_index, (label, step) in enumerate(zip(labels, steps), start=5):
        set_text(worksheet.cell(row_index, 1), label)
        set_text(worksheet.cell(row_index, 4), step[0])
        worksheet.merge_cells(start_row=row_index, start_column=5, end_row=row_index, end_column=8)
        set_text(worksheet.cell(row_index, 5), step[1])

    max_row = max(row_count + 6, 206)
    worksheet["B5"] = (
        f'=COUNTIFS(\'材料配置\'!$A$7:$A${max_row},"<>",'
        f'\'材料配置\'!$F$7:$F${max_row},"<>源文件已移除")'
    )
    worksheet["B6"] = (
        f'=COUNTIFS(\'材料配置\'!$A$7:$A${max_row},"<>",'
        f'\'材料配置\'!$F$7:$F${max_row},"<>源文件已移除",'
        f'\'材料配置\'!$D$7:$D${max_row},"是")'
    )
    worksheet["B7"] = (
        f'=COUNTIFS(\'材料配置\'!$A$7:$A${max_row},"<>",'
        f'\'材料配置\'!$F$7:$F${max_row},"<>源文件已移除",'
        f'\'材料配置\'!$D$7:$D${max_row},"否")'
    )
    worksheet["B8"] = (
        f'=COUNTIFS(\'材料配置\'!$A$7:$A${max_row},"<>",'
        f'\'材料配置\'!$E$7:$E${max_row},"<>通过")'
    )

    for row_index in range(5, 9):
        left = worksheet.cell(row_index, 1)
        left.fill = fill(COLORS["blue"])
        left.font = Font(bold=True, color=COLORS["white"])
        left.alignment = Alignment(horizontal="center", vertical="center")
        left.border = THIN_BORDER
        value = worksheet.cell(row_index, 2)
        value.font = Font(bold=True, color=COLORS["navy"], size=15)
        value.alignment = Alignment(horizontal="center", vertical="center")
        value.border = THIN_BORDER
        value.number_format = "0"
        step_label = worksheet.cell(row_index, 4)
        step_label.fill = fill(COLORS["pale_gray"])
        step_label.font = Font(bold=True, color=COLORS["navy"])
        step_label.alignment = Alignment(horizontal="center", vertical="center")
        step_label.border = THIN_BORDER
        apply_cells(
            worksheet,
            f"E{row_index}:H{row_index}",
            lambda cell: style_grid_cell(cell),
        )

    worksheet.merge_cells("A11:H11")
    set_text(worksheet["A11"], "你需要修改的只有 3 列")
    worksheet["A11"].fill = fill(COLORS["green"])
    worksheet["A11"].font = Font(bold=True, color=COLORS["white"], size=12)
    instructions = [
        ("绘图名称（可修改）", "用于所有图题、图例和输出表；请填写完整材料名称。"),
        ("进入总结图集（是/否）", "选择“否”后仍保留分析数据，但不会出现在总结图与 PDF 明细页。"),
        ("用户备注（可填写）", "可记录命名依据、批次、排除原因等；不会参与计算。"),
        ("保存要求", "修改后请保存并关闭 Excel，再执行步骤 2。"),
    ]
    for row_index, (label, explanation) in enumerate(instructions, start=12):
        set_text(worksheet.cell(row_index, 1), label)
        worksheet.cell(row_index, 1).fill = fill(COLORS["yellow"])
        worksheet.cell(row_index, 1).font = Font(bold=True, color=COLORS["text"])
        worksheet.cell(row_index, 1).border = THIN_BORDER
        worksheet.merge_cells(start_row=row_index, start_column=2, end_row=row_index, end_column=8)
        set_text(worksheet.cell(row_index, 2), explanation)
        apply_cells(
            worksheet,
            f"B{row_index}:H{row_index}",
            lambda cell: style_grid_cell(cell),
        )

    worksheet.merge_cells("A17:H17")
    set_text(worksheet["A17"], "本次扫描信息")
    worksheet["A17"].fill = fill(COLORS["pale_gray"])
    worksheet["A17"].font = Font(bold=True, color=COLORS["navy"])
    scan_rows = [
        ("扫描时间", snapshot.get("generated_at", "")),
        ("扫描位置", snapshot.get("source_root", "")),
        ("数据快照指纹（勿改）", snapshot.get("dataset_fingerprint", "")),
    ]
    for row_index, (label, value) in enumerate(scan_rows, start=18):
        set_text(worksheet.cell(row_index, 1), label)
        worksheet.cell(row_index, 1).fill = fill(COLORS["pale_gray"])
        worksheet.cell(row_index, 1).font = Font(bold=True)
        worksheet.cell(row_index, 1).border = THIN_BORDER
        worksheet.merge_cells(start_row=row_index, start_column=2, end_row=row_index, end_column=8)
        set_text(worksheet.cell(row_index, 2), value)
        apply_cells(
            worksheet,
            f"B{row_index}:H{row_index}",
            lambda cell: style_grid_cell(cell),
        )

    set_column_widths(
        worksheet,
        {"A": 24, "B": 13, "C": 3, "D": 14, "E": 18, "F": 18, "G": 18, "H": 18},
    )
    worksheet.freeze_panes = "A4"


def build_material_sheet(workbook: Workbook, rows: list[dict[str, Any]]) -> None:
    worksheet = workbook.create_sheet("材料配置")
    worksheet.sheet_view.showGridLines = False
    apply_title_band(worksheet, "A1:P1", "材料配置｜黄色列可修改")
    worksheet.merge_cells("A2:P2")
    set_text(
        worksheet["A2"],
        "刷新配置表会按“材料键”继承绘图名称、入图选择和备注；请不要修改材料键、数据状态和数据指纹。",
    )
    worksheet["A2"].fill = fill(COLORS["pale_blue"])
    worksheet["A2"].font = Font(color=COLORS["navy"], bold=True)
    worksheet["A2"].alignment = Alignment(wrap_text=True)

    current_count = sum(row["status"] != "源文件已移除" for row in rows)
    included_count = sum(
        row["status"] != "源文件已移除" and row["include"] == "是" for row in rows
    )
    excluded_count = sum(
        row["status"] != "源文件已移除" and row["include"] == "否" for row in rows
    )
    changed_count = sum(row["status"] in {"新增", "数据已更新"} for row in rows)
    summary = [
        "当前材料数",
        current_count,
        "默认入图",
        included_count,
        "已排除",
        excluded_count,
        "本次新增或更新",
        changed_count,
    ]
    for column_index, value in enumerate(summary, start=1):
        cell = worksheet.cell(4, column_index)
        if isinstance(value, str):
            set_text(cell, value)
            cell.fill = fill(COLORS["blue"])
            cell.font = Font(bold=True, color=COLORS["white"])
        else:
            cell.value = value
            cell.font = Font(bold=True, color=COLORS["navy"], size=13)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = THIN_BORDER

    for column_index, header in enumerate(MATERIAL_HEADERS, start=1):
        cell = worksheet.cell(6, column_index)
        set_text(cell, header)
        cell.fill = fill(COLORS["navy"])
        cell.font = Font(bold=True, color=COLORS["white"])
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER
    worksheet.row_dimensions[6].height = 34

    first_row = 7
    last_row = first_row + len(rows) - 1
    for row_index, item in enumerate(rows, start=first_row):
        text_values = {
            1: item["key"],
            2: item["auto_name"],
            3: item["plot_name"],
            4: item["include"],
            6: item["status"],
            9: item["order"],
            10: item["relative_path"],
            11: item["fingerprint"],
            12: item["notes"],
            13: item["scanned_at"],
            14: item["test_types"],
        }
        for column_index, value in text_values.items():
            set_text(worksheet.cell(row_index, column_index), value)
        worksheet.cell(row_index, 5).value = (
            f'=IF(C{row_index}="","需要填写名称",'
            f'IF(COUNTIF($C${first_row}:$C${last_row},C{row_index})>1,"绘图名称重复",'
            f'IF(OR(D{row_index}="是",D{row_index}="否"),"通过","选择值无效")))'
        )
        worksheet.cell(row_index, 7).value = item["standard_file_count"]
        worksheet.cell(row_index, 8).value = item["total_data_points"]
        worksheet.cell(row_index, 15).value = item["cathodic_current"]
        worksheet.cell(row_index, 16).value = item["recovery_current"]
        for column_index in range(1, 17):
            style_grid_cell(worksheet.cell(row_index, column_index), body=True)
        worksheet.cell(row_index, 3).fill = fill(COLORS["yellow"])
        worksheet.cell(row_index, 4).fill = fill(COLORS["yellow"])
        worksheet.cell(row_index, 12).fill = fill(COLORS["yellow"])
        worksheet.cell(row_index, 7).number_format = "0"
        worksheet.cell(row_index, 8).number_format = "0"
        worksheet.cell(row_index, 15).number_format = "0.0000"
        worksheet.cell(row_index, 16).number_format = "0.0000"

    if rows:
        validation = DataValidation(type="list", formula1='"是,否"', allow_blank=False)
        worksheet.add_data_validation(validation)
        validation.add(f"D{first_row}:D{last_row}")

        green_fill = fill(COLORS["pale_green"])
        red_fill = fill(COLORS["pale_red"])
        blue_fill = fill(COLORS["pale_blue"])
        orange_fill = fill(COLORS["pale_orange"])
        gray_fill = fill(COLORS["pale_gray"])
        worksheet.conditional_formatting.add(
            f"D{first_row}:D{last_row}",
            FormulaRule(
                formula=[f'D{first_row}="是"'],
                fill=green_fill,
                font=Font(color=COLORS["green"], bold=True),
            ),
        )
        worksheet.conditional_formatting.add(
            f"D{first_row}:D{last_row}",
            FormulaRule(
                formula=[f'D{first_row}="否"'],
                fill=red_fill,
                font=Font(color=COLORS["red"], bold=True),
            ),
        )
        worksheet.conditional_formatting.add(
            f"E{first_row}:E{last_row}",
            FormulaRule(
                formula=[f'E{first_row}="通过"'],
                fill=green_fill,
                font=Font(color=COLORS["green"]),
            ),
        )
        worksheet.conditional_formatting.add(
            f"E{first_row}:E{last_row}",
            FormulaRule(
                formula=[f'E{first_row}<>"通过"'],
                fill=red_fill,
                font=Font(color=COLORS["red"], bold=True),
            ),
        )
        for status, status_fill, color in [
            ("新增", blue_fill, COLORS["navy"]),
            ("数据已更新", orange_fill, COLORS["orange"]),
            ("源文件已移除", gray_fill, COLORS["gray"]),
        ]:
            worksheet.conditional_formatting.add(
                f"F{first_row}:F{last_row}",
                FormulaRule(
                    formula=[f'F{first_row}="{status}"'],
                    fill=status_fill,
                    font=Font(color=color, bold=status != "源文件已移除"),
                ),
            )

        table = Table(displayName="MaterialConfigTable", ref=f"A6:P{last_row}")
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        worksheet.add_table(table)

    set_column_widths(
        worksheet,
        {
            "A": 29,
            "B": 24,
            "C": 28,
            "D": 20,
            "E": 18,
            "F": 18,
            "G": 13,
            "H": 14,
            "I": 66,
            "J": 36,
            "K": 33,
            "L": 34,
            "M": 21,
            "N": 18,
            "O": 20,
            "P": 20,
        },
    )
    worksheet.freeze_panes = "C7"


def build_file_sheet(workbook: Workbook, snapshot: dict[str, Any]) -> None:
    worksheet = workbook.create_sheet("文件清单")
    worksheet.sheet_view.showGridLines = False
    apply_title_band(worksheet, "A1:O1", "本次扫描文件清单｜只读")
    worksheet.merge_cells("A2:O2")
    set_text(
        worksheet["A2"],
        "这里列出所有识别到的启停候选文件；未纳入的文件会保留排除原因，便于核对。",
    )
    worksheet["A2"].fill = fill(COLORS["pale_blue"])
    worksheet["A2"].font = Font(color=COLORS["navy"], bold=True)
    worksheet["A2"].alignment = Alignment(wrap_text=True)

    for column_index, header in enumerate(FILE_HEADERS, start=1):
        cell = worksheet.cell(4, column_index)
        set_text(cell, header)
        cell.fill = fill(COLORS["navy"])
        cell.font = Font(bold=True, color=COLORS["white"])
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER
    worksheet.row_dimensions[4].height = 34

    files = snapshot.get("files", [])
    if not isinstance(files, list):
        raise ValueError("材料快照中的 files 必须是数组。")
    for row_index, item in enumerate(files, start=5):
        if not isinstance(item, dict):
            raise ValueError("材料快照中的文件记录必须是对象。")
        included = bool(item.get("included_in_analysis"))
        status = (
            "纳入启停"
            if included
            else normalize_cell(item.get("exclusion_reason"))
            or normalize_cell(item.get("candidate_class"))
        )
        text_values = {
            1: item.get("relative_path", ""),
            2: item.get("material_relative_path", ""),
            3: item.get("file_name", ""),
            4: item.get("file_name_kind", ""),
            5: bool_zh(included),
            6: status,
            8: item.get("current_levels_a_cm2", ""),
            11: item.get("file_mtime", ""),
            12: item.get("sha256", ""),
            13: item.get("test_type_label_zh", ""),
        }
        for column_index, value in text_values.items():
            set_text(worksheet.cell(row_index, column_index), value)
        worksheet.cell(row_index, 7).value = safe_number(item.get("valid_row_count"))
        worksheet.cell(row_index, 9).value = safe_number(item.get("time_start_s"))
        worksheet.cell(row_index, 10).value = safe_number(item.get("time_end_s"))
        worksheet.cell(row_index, 14).value = safe_number(
            item.get("cathodic_current_median_a_cm2")
        )
        worksheet.cell(row_index, 15).value = safe_number(
            item.get("recovery_current_median_a_cm2")
        )
        for column_index in range(1, 16):
            style_grid_cell(worksheet.cell(row_index, column_index), body=True)
        worksheet.cell(row_index, 7).number_format = "0"
        worksheet.cell(row_index, 9).number_format = "0.000"
        worksheet.cell(row_index, 10).number_format = "0.000"
        worksheet.cell(row_index, 14).number_format = "0.0000"
        worksheet.cell(row_index, 15).number_format = "0.0000"

    if files:
        last_row = 4 + len(files)
        worksheet.conditional_formatting.add(
            f"E5:E{last_row}",
            FormulaRule(
                formula=['E5="是"'],
                fill=fill(COLORS["pale_green"]),
                font=Font(color=COLORS["green"], bold=True),
            ),
        )
        worksheet.conditional_formatting.add(
            f"E5:E{last_row}",
            FormulaRule(
                formula=['E5="否"'],
                fill=fill(COLORS["pale_gray"]),
                font=Font(color=COLORS["gray"]),
            ),
        )
        table = Table(displayName="FileInventoryTable", ref=f"A4:O{last_row}")
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        worksheet.add_table(table)

    set_column_widths(
        worksheet,
        {
            "A": 49,
            "B": 34,
            "C": 31,
            "D": 23,
            "E": 20,
            "F": 48,
            "G": 15,
            "H": 27,
            "I": 16,
            "J": 16,
            "K": 22,
            "L": 34,
            "M": 18,
            "N": 20,
            "O": 20,
        },
    )
    worksheet.freeze_panes = "C5"


def build_workbook(snapshot: dict[str, Any], rows: list[dict[str, Any]]) -> Workbook:
    workbook = Workbook()
    workbook.remove(workbook.active)
    build_usage_sheet(workbook, snapshot, len(rows))
    build_material_sheet(workbook, rows)
    build_file_sheet(workbook, snapshot)
    workbook.calculation.calcMode = "auto"
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    return workbook


def update_workbook(snapshot_path: Path, workbook_path: Path) -> dict[str, Any]:
    snapshot = load_json(snapshot_path)
    existing = load_existing_config(workbook_path)
    rows = build_material_rows(snapshot, existing)
    workbook = build_workbook(snapshot, rows)
    try:
        atomic_save_workbook(workbook, workbook_path)
    finally:
        workbook.close()
    return {
        "workbook": str(workbook_path),
        "current_materials": len(snapshot.get("materials", [])),
        "rows": len(rows),
        "preserved_rows": sum(
            item["key"] in existing["rows_by_key"] for item in rows
        ),
        "dataset_changed": (
            existing.get("dataset_fingerprint", "")
            != normalize_cell(snapshot.get("dataset_fingerprint"))
        ),
    }


def read_workbook(workbook_path: Path, output_path: Path) -> dict[str, Any]:
    workbook = load_workbook(workbook_path, data_only=False, read_only=False)
    try:
        if "使用说明" not in workbook.sheetnames or "材料配置" not in workbook.sheetnames:
            raise ValueError("材料配置表缺少标准工作表，请先执行步骤 1 更新配置表。")
        dataset_fingerprint = normalize_cell(workbook["使用说明"]["B20"].value)
        worksheet = workbook["材料配置"]
        header_row = find_material_header_row(worksheet)
        if header_row is None:
            raise ValueError("材料配置表缺少标准表头，请先执行步骤 1 更新配置表。")

        rows: list[dict[str, Any]] = []
        for row_index in range(header_row + 1, worksheet.max_row + 1):
            key = normalize_cell(worksheet.cell(row_index, 1).value)
            if not key:
                continue
            status = normalize_cell(worksheet.cell(row_index, 6).value)
            rows.append(
                {
                    "key": key,
                    "auto_name": normalize_cell(worksheet.cell(row_index, 2).value),
                    "plot_name": normalize_cell(worksheet.cell(row_index, 3).value),
                    "include_in_summary_atlas": normalize_cell(
                        worksheet.cell(row_index, 4).value
                    ),
                    "status": status,
                    "fingerprint": normalize_cell(worksheet.cell(row_index, 11).value),
                    "notes": normalize_cell(worksheet.cell(row_index, 12).value),
                    "source_removed": status == "源文件已移除",
                }
            )
    finally:
        workbook.close()

    current_rows = [item for item in rows if not item["source_removed"]]
    errors: list[str] = []
    seen_names: dict[str, str] = {}
    for item in current_rows:
        if not item["plot_name"]:
            errors.append(f'{item["key"]}：绘图名称为空')
        if item["include_in_summary_atlas"] not in {"是", "否"}:
            errors.append(f'{item["key"]}：进入总结图集必须选择“是”或“否”')
        if item["plot_name"]:
            previous_key = seen_names.get(item["plot_name"])
            if previous_key:
                errors.append(
                    f'{item["key"]}：绘图名称“{item["plot_name"]}”与 {previous_key} 重复'
                )
            else:
                seen_names[item["plot_name"]] = item["key"]
    selected = sum(item["include_in_summary_atlas"] == "是" for item in current_rows)
    if selected == 0:
        errors.append("至少需要选择 1 种材料进入总结图集")
    if errors:
        raise ValueError("配置表需要修改：\n- " + "\n- ".join(errors))

    payload = {
        "dataset_fingerprint": dataset_fingerprint,
        "materials": [
            {
                **item,
                "include_in_summary_atlas": item["include_in_summary_atlas"] == "是",
            }
            for item in current_rows
        ],
    }
    atomic_write_json(output_path, payload)
    return {
        "output": str(output_path),
        "materials": len(current_rows),
        "selected": selected,
    }


def apply_json_config(
    snapshot_path: Path,
    config_path: Path,
    workbook_path: Path,
) -> dict[str, Any]:
    snapshot = load_json(snapshot_path)
    config = load_json(config_path)
    if config.get("dataset_fingerprint") != snapshot.get("dataset_fingerprint"):
        raise ValueError("网页配置对应的数据快照已过期，请先更新数据。")
    config_materials = config.get("materials")
    if not isinstance(config_materials, list):
        raise ValueError("网页材料配置缺少 materials 数组。")
    snapshot_materials = snapshot.get("materials")
    if not isinstance(snapshot_materials, list):
        raise ValueError("材料快照缺少 materials 数组。")

    current_by_key = {
        normalize_cell(item.get("key")): item
        for item in snapshot_materials
        if isinstance(item, dict)
    }
    rows_by_key: dict[str, dict[str, str]] = {}
    seen_names: dict[str, str] = {}
    errors: list[str] = []
    selected = 0
    for raw_item in config_materials:
        if not isinstance(raw_item, dict):
            errors.append("(空材料键)：材料键与当前快照不一致")
            continue
        key = normalize_cell(raw_item.get("key"))
        material = current_by_key.get(key)
        if material is None or key in rows_by_key:
            errors.append(f'{key or "(空材料键)"}：材料键与当前快照不一致')
            continue
        plot_name = normalize_cell(raw_item.get("plot_name"))
        include = raw_item.get("include_in_summary_atlas")
        fingerprint = normalize_cell(raw_item.get("fingerprint"))
        notes = normalize_cell(raw_item.get("notes"))
        if not plot_name:
            errors.append(f"{key}：绘图名称为空")
        if not isinstance(include, bool):
            errors.append(f"{key}：进入总结图集必须是布尔值")
        if fingerprint != normalize_cell(material.get("fingerprint")):
            errors.append(f"{key}：源数据已更新")
        if plot_name:
            previous_key = seen_names.get(plot_name)
            if previous_key:
                errors.append(f"{key}：绘图名称“{plot_name}”与 {previous_key} 重复")
            else:
                seen_names[plot_name] = key
        if include is True:
            selected += 1
        rows_by_key[key] = {
            "key": key,
            "auto_name": normalize_cell(raw_item.get("auto_name"))
            or normalize_cell(material.get("auto_name")),
            "plot_name": plot_name,
            "include": "是" if include is True else "否",
            "status": normalize_cell(raw_item.get("status")),
            "fingerprint": fingerprint,
            "notes": notes,
        }
    if len(rows_by_key) != len(current_by_key):
        errors.append("网页配置中的材料集合不完整")
    if selected == 0:
        errors.append("至少需要选择 1 种材料进入总结图集")
    if errors:
        raise ValueError("网页材料配置无法应用：\n- " + "\n- ".join(errors))

    rows = build_material_rows(
        snapshot,
        {
            "rows_by_key": rows_by_key,
            "dataset_fingerprint": normalize_cell(snapshot.get("dataset_fingerprint")),
        },
    )
    workbook = build_workbook(snapshot, rows)
    try:
        atomic_save_workbook(workbook, workbook_path)
    finally:
        workbook.close()
    return {
        "workbook": str(workbook_path),
        "materials": len(rows),
        "selected": selected,
        "dataset_fingerprint": normalize_cell(snapshot.get("dataset_fingerprint")),
    }


def emit(payload: dict[str, Any]) -> None:
    sys.stdout.write(json.dumps(payload, ensure_ascii=False, indent=2) + "\n")


def main(arguments: list[str]) -> int:
    if len(arguments) == 3 and arguments[0] == "update":
        emit(update_workbook(Path(arguments[1]), Path(arguments[2])))
        return 0
    if len(arguments) == 3 and arguments[0] == "read":
        emit(read_workbook(Path(arguments[1]), Path(arguments[2])))
        return 0
    if len(arguments) == 4 and arguments[0] == "apply-json":
        emit(
            apply_json_config(
                Path(arguments[1]),
                Path(arguments[2]),
                Path(arguments[3]),
            )
        )
        return 0
    raise ValueError(
        "用法：material_config_workbook.py update <snapshot.json> <workbook.xlsx> | "
        "read <workbook.xlsx> <output.json> | "
        "apply-json <snapshot.json> <config.json> <workbook.xlsx>"
    )


if __name__ == "__main__":
    try:
        raise SystemExit(main(sys.argv[1:]))
    except SystemExit:
        raise
    except Exception as error:
        print(str(error), file=sys.stderr)
        raise SystemExit(1) from None

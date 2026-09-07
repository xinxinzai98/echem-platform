"""Build highlighted raw/processed Excel and PDF from a fixed analysis generation."""
from __future__ import annotations

import contextlib
import csv
import datetime as dt
import errno
import io
import math
import os
import re
import shutil
import sqlite3
import statistics
import tempfile
import uuid
from collections import OrderedDict
from pathlib import Path
from typing import Any, Iterable
from .start_stop_resources import ScratchSpaceError
from .start_stop_contracts import StartStopWorkspaceError, _bool, _integer


class ChartExportMixin:
    @staticmethod
    def _excel_safe_text(value: Any) -> str:
        text = str(value or "")
        if text.startswith(("=", "+", "-", "@", "\t", "\r")):
            return "'" + text
        return text

    @staticmethod
    def _chart_variant_label(variant: str) -> str:
        return "端点处理（水位补偿）" if variant == "water" else "端点处理（未水位补偿）"

    def _raw_export_context(self, context: dict[str, Any]) -> dict[str, Any]:
        snapshot_path = self._path(self.SNAPSHOT_NAME, require=True)
        snapshot_signature = self._file_cache_signature(snapshot_path)
        snapshot = self._read_json(self.SNAPSHOT_NAME)
        raw_files = snapshot.get("files")
        if not isinstance(raw_files, list):
            raise StartStopWorkspaceError(
                "当前分析缺少原始文件快照，请重新绘图后再导出。",
                409,
            )
        snapshot_files: dict[str, dict[str, Any]] = {}
        for item in raw_files:
            if not isinstance(item, dict):
                continue
            path = str(item.get("relative_path") or "")
            if path:
                snapshot_files[path] = item

        identities: OrderedDict[tuple[str, str, int], dict[str, Any]] = OrderedDict()
        raw_series: list[dict[str, Any]] = []
        for series_id in context["requested"]:
            series_meta = context["known"][series_id]
            source_files = series_meta.get("source_files")
            if not isinstance(source_files, list) or not source_files:
                raise StartStopWorkspaceError(
                    "当前高亮序列缺少原始文件清单，请重新绘图后再导出。",
                    409,
                )
            series_files: list[dict[str, Any]] = []
            for source_order, source_file in enumerate(source_files, start=1):
                repository_path = str(source_file or "")
                snapshot_row = snapshot_files.get(repository_path)
                if not isinstance(snapshot_row, dict) or not _bool(
                    snapshot_row.get("included_in_analysis")
                ):
                    raise StartStopWorkspaceError(
                        "当前高亮序列的原始文件不在已发布分析快照中。",
                        409,
                    )
                snapshot_series_id = str(
                    snapshot_row.get("analysis_series_id") or ""
                )
                if snapshot_series_id and snapshot_series_id != series_id:
                    raise StartStopWorkspaceError(
                        "当前高亮序列与原始文件快照不一致，请重新绘图后再导出。",
                        409,
                    )
                sha256 = str(snapshot_row.get("sha256") or "").strip().casefold()
                if re.fullmatch(r"[0-9a-f]{64}", sha256) is None:
                    raise StartStopWorkspaceError(
                        "当前原始文件快照缺少有效校验值，请重新绘图后再导出。",
                        409,
                    )
                size_bytes = _integer(snapshot_row.get("file_size_bytes"), -1)
                if size_bytes < 0:
                    raise StartStopWorkspaceError(
                        "当前原始文件快照缺少文件大小，请重新绘图后再导出。",
                        409,
                    )
                if size_bytes > self.MAX_RAW_EXPORT_FILE_BYTES:
                    raise StartStopWorkspaceError(
                        "当前高亮材料包含过大的单个原始文件，请减少高亮材料后再导出。",
                        413,
                    )
                identity = (repository_path, sha256, size_bytes)
                identities.setdefault(
                    identity,
                    {
                        "repository_path": repository_path,
                        "sha256": sha256,
                        "size_bytes": size_bytes,
                    },
                )
                series_files.append(
                    {
                        "identity": identity,
                        "repository_path": repository_path,
                        "segment_index": _integer(
                            snapshot_row.get("segment_index"),
                            source_order,
                        ),
                        "snapshot_raw_row_count": _integer(
                            snapshot_row.get("raw_row_count")
                        ),
                        "snapshot_valid_row_count": _integer(
                            snapshot_row.get("valid_row_count")
                        ),
                    }
                )
            raw_series.append(
                {
                    "series_id": series_id,
                    "series_meta": series_meta,
                    "files": series_files,
                }
            )

        manifest_reader = getattr(
            self.source_database,
            "source_export_manifest",
            None,
        )
        if not callable(manifest_reader):
            raise StartStopWorkspaceError(
                "原始文件数据库读取组件当前不可用。",
                503,
            )
        try:
            manifest = manifest_reader(
                list(identities.values()),
                maximum_files=self.MAX_RAW_EXPORT_SOURCE_FILES,
                maximum_total_bytes=self.MAX_RAW_EXPORT_TOTAL_BYTES,
            )
        except OverflowError as exc:
            raise StartStopWorkspaceError(
                "当前高亮材料的原始文件过大，请减少高亮材料后再导出。",
                413,
            ) from exc
        except KeyError as exc:
            raise StartStopWorkspaceError(
                "已发布分析对应的原始文件版本不在数据库中。",
                409,
            ) from exc
        except (OSError, RuntimeError, ValueError, sqlite3.Error) as exc:
            raise StartStopWorkspaceError(
                "无法核对当前高亮材料的原始文件。",
                500,
            ) from exc
        records = {
            (
                str(item.get("repository_path") or ""),
                str(item.get("sha256") or ""),
                _integer(item.get("size_bytes"), -1),
            ): item
            for item in manifest
            if isinstance(item, dict)
        }
        if set(records) != set(identities):
            raise StartStopWorkspaceError(
                "原始文件数据库回读结果与已发布分析不一致。",
                409,
            )
        for series in raw_series:
            for item in series["files"]:
                item["database_record"] = records[item["identity"]]
        return {
            "dataset_fingerprint": str(snapshot.get("dataset_fingerprint") or ""),
            "snapshot_path": snapshot_path,
            "snapshot_signature": snapshot_signature,
            "generation": context["generation"],
            "series": raw_series,
            "file_count": len(identities),
            "total_bytes": sum(item[2] for item in identities),
        }

    @staticmethod
    def _decode_raw_source(content: bytes) -> tuple[str, str]:
        for encoding in ("utf-8-sig", "gb18030", "utf-16"):
            try:
                text = content.decode(encoding)
            except UnicodeError:
                continue
            if "\x00" not in text[:4096]:
                return text, encoding
        return content.decode("utf-8-sig", errors="replace"), "utf-8-sig（替换无效字节）"

    @staticmethod
    def _raw_source_cells(line: str, delimiter: str) -> list[str]:
        try:
            return next(csv.reader([line], delimiter=delimiter))
        except (csv.Error, StopIteration):
            return [line]

    @staticmethod
    def _raw_source_number(value: Any) -> float | None:
        try:
            parsed = float(str(value).strip())
        except (TypeError, ValueError):
            return None
        return parsed if math.isfinite(parsed) else None

    def _profile_raw_source(self, content: bytes) -> dict[str, Any]:
        text, encoding = self._decode_raw_source(content)
        header_index = -1
        header_line = ""
        preamble_lines: list[str] = []
        for index, line in enumerate(io.StringIO(text)):
            if line.lstrip("\ufeff").startswith("E(V)"):
                header_index = index
                header_line = line.rstrip("\r\n")
                break
            preamble_lines.append(line.rstrip("\r\n"))
            if index >= 50:
                break
        if header_index < 0:
            raise StartStopWorkspaceError(
                "原始文件未找到 E(V) 数据表头。",
                422,
            )
        delimiter = "\t" if "\t" in header_line else "," if "," in header_line else "\t"
        header_cells = self._raw_source_cells(header_line, delimiter)
        valid_times: list[float] = []
        data_row_count = 0
        valid_row_count = 0
        blank_row_count = 0
        stream = io.StringIO(text)
        for _ in range(header_index + 1):
            next(stream, "")
        for line in stream:
            raw_line = line.rstrip("\r\n")
            if not raw_line:
                blank_row_count += 1
                continue
            data_row_count += 1
            cells = self._raw_source_cells(raw_line, delimiter)
            values = [
                self._raw_source_number(cells[index]) if index < len(cells) else None
                for index in range(3)
            ]
            if all(value is not None for value in values):
                valid_row_count += 1
                valid_times.append(float(values[2]))
        positive_steps = [
            following - previous
            for previous, following in zip(valid_times, valid_times[1:])
            if following - previous > 0
        ]
        sample_interval_s = (
            float(statistics.median(positive_steps)) if positive_steps else 0.0
        )
        return {
            "text": text,
            "encoding": encoding,
            "header_index": header_index,
            "header_line": header_line,
            "header_cells": header_cells,
            "preamble_text": "\n".join(preamble_lines),
            "delimiter": delimiter,
            "data_row_count": data_row_count,
            "valid_row_count": valid_row_count,
            "parse_error_rows": data_row_count - valid_row_count,
            "blank_row_count": blank_row_count,
            "sample_interval_s": sample_interval_s,
            "time_start_s": valid_times[0] if valid_times else None,
            "time_end_s": valid_times[-1] if valid_times else None,
        }

    def _iter_profiled_raw_rows(
        self,
        profile: dict[str, Any],
    ) -> Iterable[tuple[int, str, list[str], tuple[float | None, float | None, float | None]]]:
        stream = io.StringIO(profile["text"])
        for _ in range(int(profile["header_index"]) + 1):
            next(stream, "")
        first_line_number = int(profile["header_index"]) + 2
        for line_number, line in enumerate(stream, start=first_line_number):
            raw_line = line.rstrip("\r\n")
            if not raw_line:
                continue
            cells = self._raw_source_cells(raw_line, profile["delimiter"])
            values = tuple(
                self._raw_source_number(cells[index]) if index < len(cells) else None
                for index in range(3)
            )
            yield line_number, raw_line, cells, values

    def _build_highlight_excel(
        self,
        *,
        context: dict[str, Any],
        raw_context: dict[str, Any],
        grouped: dict[tuple[str, str], list[dict[str, Any]]],
        names: dict[str, str],
        generated_at: dt.datetime,
        workbook_lifecycle: list[Any] | None = None,
    ) -> bytes:
        try:
            from openpyxl import Workbook
            from openpyxl.cell import WriteOnlyCell
            from openpyxl.styles import Font, PatternFill
        except ImportError as exc:
            raise StartStopWorkspaceError(
                "Excel 导出组件当前不可用。",
                503,
            ) from exc

        workbook = Workbook(write_only=True)
        if workbook_lifecycle is not None:
            workbook_lifecycle.append(workbook)
        workbook.properties.creator = "Start-stop Studio"
        workbook.properties.title = "启停分析当前高亮材料原始与处理数据"
        workbook.properties.subject = "完整原始测量点、接续时间与启停处理结果"
        header_fill = PatternFill("solid", fgColor="0B7F77")
        header_font = Font(color="FFFFFF", bold=True)

        def bounded_text(
            value: Any,
            *,
            label: str,
            formula_safe: bool = True,
        ) -> str:
            text = str(value or "")
            if formula_safe:
                text = self._excel_safe_text(text)
            if len(text) > self.MAX_EXCEL_CELL_TEXT:
                raise StartStopWorkspaceError(
                    f"{label}超过 Excel 单元格可保存长度，无法无损导出。",
                    413,
                )
            if any(
                ord(character) < 32 and character not in {"\t", "\n", "\r"}
                for character in text
            ):
                raise StartStopWorkspaceError(
                    f"{label}包含 Excel 无法保存的控制字符。",
                    422,
                )
            return text

        def append_header(sheet: Any, values: Iterable[Any]) -> None:
            cells = []
            for value in values:
                cell = WriteOnlyCell(
                    sheet,
                    value=bounded_text(value, label="工作表表头"),
                )
                cell.fill = header_fill
                cell.font = header_font
                cells.append(cell)
            sheet.append(cells)

        def append_formatted_row(
            sheet: Any,
            values: Iterable[Any],
            number_formats: Iterable[str | None],
            *,
            force_text_columns: Iterable[int] = (),
        ) -> None:
            cells = []
            forced = set(force_text_columns)
            for column_index, (value, number_format) in enumerate(
                zip(values, number_formats)
            ):
                cell = WriteOnlyCell(sheet, value=value)
                if column_index in forced and value is not None:
                    cell.data_type = "s"
                if number_format is not None and value is not None:
                    cell.number_format = number_format
                cells.append(cell)
            sheet.append(cells)

        mode_labels = {
            "raw": "端点处理（未水位补偿）",
            "water": "端点处理（水位补偿）",
            "compare": "两种端点处理结果对照",
        }
        first_series = context["known"][context["requested"][0]]
        info = workbook.create_sheet("导出说明")
        info.sheet_view.showGridLines = False
        info.column_dimensions["A"].width = 20
        info.column_dimensions["B"].width = 72
        append_header(info, ("字段", "内容"))
        for key, value in (
            ("生成时间", generated_at.isoformat().replace("+00:00", "Z")),
            ("数据来源", "Start-stop Studio 当前已发布分析结果"),
            ("启停工步", first_series.get("work_step_label") or context["work_step_key"]),
            ("分析指标", context["metric_spec"]["label"]),
            ("单位", context["metric_spec"]["unit"]),
            ("横轴", "累计时间 / h" if context["overview"] or context["x_axis"] == "time" else "循环数"),
            ("数据模式", mode_labels[context["mode"]]),
            ("高亮序列数", len(context["requested"])),
            ("电位基准", "Hg/HgO 原始标尺；不进行 RHE 换算"),
            (
                "完整原始文件",
                f"{raw_context['file_count']} 个文件，"
                f"{raw_context['total_bytes']} 字节；均按当前分析快照的 SHA-256 回读",
            ),
            (
                "原始数据接续",
                "同一材料按已发布 source_files 顺序接续；每个文件先将原始时间减去首个有效时间，再累加前段含末端采样间隔的时长",
            ),
            (
                "原始数据未处理项",
                "不提取阶段末点、不筛选循环、不做水位补偿、不做 RHE 换算；仅新增接续时间、来源与解析状态列",
            ),
            (
                "原始数据可追溯性",
                "每行保留来源文件、接续段、原文件行号和原始整行；文件前置信息、表头、SHA-256 与数据库版本见“原始文件索引”",
            ),
            ("阴极末端规则", "阶段最后 1 s 中位数；ADT 使用末点"),
            ("异常规则", "最低点 < 15 s 为异常，>= 15 s 为正常"),
        ):
            info.append(
                (
                    bounded_text(key, label="导出说明"),
                    bounded_text(value, label="导出说明"),
                )
            )

        index_sheet = workbook.create_sheet("曲线索引")
        index_sheet.sheet_view.showGridLines = False
        for column, width in {
            "A": 8, "B": 15, "C": 40, "D": 24, "E": 14, "F": 12, "G": 42,
        }.items():
            index_sheet.column_dimensions[column].width = width
        append_header(
            index_sheet,
            ("编号", "数据页", "材料名称", "系列 ID", "数据模式", "点数", "材料键"),
        )
        index_sheet.freeze_panes = "A2"

        raw_index_sheet = workbook.create_sheet("原始文件索引")
        raw_index_sheet.sheet_view.showGridLines = False
        for column, width in {
            "A": 10,
            "B": 18,
            "C": 14,
            "D": 14,
            "E": 40,
            "F": 24,
            "G": 12,
            "H": 10,
            "I": 56,
            "J": 12,
            "K": 68,
            "L": 14,
            "M": 16,
            "N": 56,
            "O": 56,
            "P": 14,
            "Q": 14,
            "R": 14,
            "S": 12,
            "T": 16,
            "U": 16,
            "V": 16,
            "W": 16,
            "X": 16,
            "Y": 24,
            "Z": 24,
        }.items():
            raw_index_sheet.column_dimensions[column].width = width
        append_header(
            raw_index_sheet,
            (
                "材料序号",
                "完整原始数据页",
                "页内起始行",
                "页内结束行",
                "材料名称",
                "系列 ID",
                "文件顺序",
                "接续段",
                "来源文件",
                "数据库版本",
                "SHA-256",
                "文件字节",
                "字符编码",
                "文件前置信息",
                "原始表头",
                "原始数据行",
                "有效数值行",
                "解析失败行",
                "空白行",
                "采样间隔 / s",
                "原始起点 / s",
                "原始终点 / s",
                "接续起点 / s",
                "接续终点 / s",
                "来源修改时间",
                "数据库收录时间",
            ),
        )
        raw_index_sheet.freeze_panes = "A2"

        sheet_number = 0
        for series_id in context["requested"]:
            series_meta = context["known"][series_id]
            name = names.get(series_id) or str(
                series_meta.get("series_display_name") or series_id
            )
            for variant in context["variants"]:
                points = grouped[(series_id, variant)]
                if not points:
                    continue
                sheet_number += 1
                variant_label = self._chart_variant_label(variant)
                sheet_name = (
                    f"{sheet_number:02d}_端点补偿"
                    if variant == "water"
                    else f"{sheet_number:02d}_端点原始"
                )
                index_sheet.append(
                    (
                        sheet_number,
                        sheet_name,
                        bounded_text(name, label="材料名称"),
                        bounded_text(series_id, label="系列 ID"),
                        variant_label,
                        len(points),
                        bounded_text(
                            series_meta.get("material_relative_path"),
                            label="材料键",
                        ),
                    )
                )
                sheet = workbook.create_sheet(sheet_name)
                sheet.sheet_view.showGridLines = False
                for column, width in {
                    "A": 10, "B": 18, "C": 28, "D": 12, "E": 12, "F": 55, "G": 12,
                }.items():
                    sheet.column_dimensions[column].width = width
                x_label = (
                    "累计时间 / h"
                    if context["overview"] or context["x_axis"] == "time"
                    else "循环数"
                )
                y_label = f"{context['metric_spec']['label']} / {context['metric_spec']['unit']}"
                append_header(
                    sheet,
                    ("点序", x_label, y_label, "循环", "异常状态", "来源文件", "接续段"),
                )
                for point_index, point in enumerate(points, start=1):
                    status = {"normal": "正常", "abnormal": "异常"}.get(
                        str(point.get("status") or ""),
                        "",
                    )
                    values = (
                        point_index,
                        point["x"],
                        point["y"],
                        point.get("cycle") or None,
                        status,
                        bounded_text(
                            point.get("source_file"),
                            label="处理结果来源文件",
                        ),
                        point.get("segment_index") or None,
                    )
                    formats = (
                        "0",
                        "0.000000" if context["overview"] or context["x_axis"] == "time" else "0",
                        "0.000000",
                        "0",
                        None,
                        None,
                        "0",
                    )
                    append_formatted_row(sheet, values, formats)
                sheet.freeze_panes = "A2"
                sheet.auto_filter.ref = f"A1:G{len(points) + 1}"

        if sheet_number == 0:
            raise StartStopWorkspaceError("当前高亮曲线没有可导出的处理数据。")

        content_reader = getattr(
            self.source_database,
            "read_source_export_content",
            None,
        )
        if not callable(content_reader):
            raise StartStopWorkspaceError(
                "原始文件数据库读取组件当前不可用。",
                503,
            )

        raw_sheet_infos: list[dict[str, Any]] = []
        raw_index_rows: list[tuple[tuple[Any, ...], tuple[str | None, ...]]] = []
        raw_data_rows = 0
        for series_index, raw_series in enumerate(raw_context["series"], start=1):
            series_id = str(raw_series["series_id"])
            series_meta = raw_series["series_meta"]
            name = names.get(series_id) or str(
                series_meta.get("series_display_name") or series_id
            )
            continued_offset_s = 0.0
            raw_row_sequence = 0
            chunk_index = 0
            current_sheet_info: dict[str, Any] | None = None

            def start_raw_sheet() -> dict[str, Any]:
                nonlocal chunk_index
                chunk_index += 1
                sheet_name = (
                    f"R{series_index:02d}_全量接续"
                    if chunk_index == 1
                    else f"R{series_index:02d}_接续{chunk_index:02d}"
                )
                sheet = workbook.create_sheet(sheet_name)
                sheet.sheet_view.showGridLines = False
                for column, width in {
                    "A": 12,
                    "B": 18,
                    "C": 18,
                    "D": 12,
                    "E": 10,
                    "F": 56,
                    "G": 14,
                    "H": 22,
                    "I": 22,
                    "J": 20,
                    "K": 80,
                    "L": 24,
                }.items():
                    sheet.column_dimensions[column].width = width
                append_header(
                    sheet,
                    (
                        "原始行序",
                        "接续时间 / s",
                        "接续时间 / h",
                        "文件顺序",
                        "接续段",
                        "来源文件",
                        "原文件行号",
                        "原始电位 E(V) / V vs Hg/HgO",
                        "原始电流 i(A/cm²)",
                        "原始时间 T(s)",
                        "原始整行",
                        "解析状态",
                    ),
                )
                sheet.freeze_panes = "A2"
                result = {"sheet": sheet, "name": sheet_name, "rows": 0}
                raw_sheet_infos.append(result)
                return result

            for source_order, source in enumerate(raw_series["files"], start=1):
                database_record = source["database_record"]
                try:
                    content = content_reader(
                        source_version_id=int(
                            database_record["source_version_id"]
                        ),
                        expected_sha256=str(database_record["sha256"]),
                        expected_size_bytes=int(database_record["size_bytes"]),
                        maximum_file_bytes=self.MAX_RAW_EXPORT_FILE_BYTES,
                    )
                except OverflowError as exc:
                    raise StartStopWorkspaceError(
                        "当前高亮材料的单个原始文件过大。",
                        413,
                    ) from exc
                except (KeyError, ValueError) as exc:
                    raise StartStopWorkspaceError(
                        "已发布分析对应的原始文件版本无法回读。",
                        409,
                    ) from exc
                except (OSError, RuntimeError, sqlite3.Error) as exc:
                    raise StartStopWorkspaceError(
                        "读取当前高亮材料的原始文件失败。",
                        500,
                    ) from exc

                profile = self._profile_raw_source(content)
                del content
                expected_raw_rows = int(source["snapshot_raw_row_count"])
                expected_valid_rows = int(source["snapshot_valid_row_count"])
                if (
                    expected_raw_rows > 0
                    and profile["data_row_count"] != expected_raw_rows
                ) or (
                    expected_valid_rows > 0
                    and profile["valid_row_count"] != expected_valid_rows
                ):
                    raise StartStopWorkspaceError(
                        "原始文件回读行数与当前分析快照不一致。",
                        409,
                    )
                if profile["valid_row_count"] <= 0:
                    raise StartStopWorkspaceError(
                        "当前分析使用的原始文件没有可接续的有效数值点。",
                        422,
                    )
                sample_interval_s = float(profile["sample_interval_s"])
                original_start_s = float(profile["time_start_s"])
                original_end_s = float(profile["time_end_s"])
                if not math.isfinite(sample_interval_s) or sample_interval_s <= 0:
                    raise StartStopWorkspaceError(
                        "当前分析使用的原始文件无法确定正采样间隔。",
                        422,
                    )
                continued_start_s = continued_offset_s
                continued_end_s = (
                    continued_offset_s + original_end_s - original_start_s
                )
                inclusive_duration_s = (
                    original_end_s - original_start_s + sample_interval_s
                )
                file_parts: list[dict[str, Any]] = []
                current_part: dict[str, Any] | None = None

                for line_number, raw_line, _cells, values in self._iter_profiled_raw_rows(
                    profile
                ):
                    if (
                        current_sheet_info is None
                        or int(current_sheet_info["rows"])
                        >= self.MAX_EXCEL_DATA_ROWS_PER_SHEET
                    ):
                        current_sheet_info = start_raw_sheet()
                    excel_row = int(current_sheet_info["rows"]) + 2
                    if (
                        current_part is None
                        or current_part["sheet_name"] != current_sheet_info["name"]
                    ):
                        if current_part is not None:
                            file_parts.append(current_part)
                        current_part = {
                            "sheet_name": current_sheet_info["name"],
                            "row_start": excel_row,
                            "row_end": excel_row,
                        }
                    else:
                        current_part["row_end"] = excel_row

                    raw_row_sequence += 1
                    valid_point = all(value is not None for value in values)
                    potential_v, current_a_cm2, original_time_s = values
                    continued_time_s = (
                        continued_offset_s
                        + float(original_time_s)
                        - original_start_s
                        if valid_point
                        else None
                    )
                    append_formatted_row(
                        current_sheet_info["sheet"],
                        (
                            raw_row_sequence,
                            continued_time_s,
                            (
                                continued_time_s / 3600.0
                                if continued_time_s is not None
                                else None
                            ),
                            source_order,
                            int(source["segment_index"]),
                            bounded_text(
                                source["repository_path"],
                                label="原始数据来源文件",
                            ),
                            line_number,
                            potential_v,
                            current_a_cm2,
                            original_time_s,
                            bounded_text(
                                raw_line,
                                label="原始数据整行",
                                formula_safe=False,
                            ),
                            (
                                "有效原始点"
                                if valid_point
                                else "数值解析失败（原行已保留）"
                            ),
                        ),
                        (
                            "0",
                            "0.000000",
                            "0.000000000",
                            "0",
                            "0",
                            None,
                            "0",
                            "0.000000000",
                            "0.000000000",
                            "0.000000",
                            None,
                            None,
                        ),
                        force_text_columns=(10,),
                    )
                    current_sheet_info["rows"] = int(current_sheet_info["rows"]) + 1
                    raw_data_rows += 1

                if current_part is not None:
                    file_parts.append(current_part)
                if not file_parts:
                    raise StartStopWorkspaceError(
                        "当前分析使用的原始文件没有可导出的数据行。",
                        422,
                    )

                for file_part in file_parts:
                    raw_index_rows.append(
                        (
                            (
                                series_index,
                                file_part["sheet_name"],
                                file_part["row_start"],
                                file_part["row_end"],
                                bounded_text(name, label="材料名称"),
                                bounded_text(series_id, label="系列 ID"),
                                source_order,
                                int(source["segment_index"]),
                                bounded_text(
                                    source["repository_path"],
                                    label="原始文件路径",
                                ),
                                int(database_record["version_number"]),
                                bounded_text(
                                    database_record["sha256"],
                                    label="原始文件校验值",
                                ),
                                int(database_record["size_bytes"]),
                                bounded_text(
                                    profile["encoding"],
                                    label="原始文件编码",
                                ),
                                bounded_text(
                                    profile["preamble_text"],
                                    label="原始文件前置信息",
                                    formula_safe=False,
                                ),
                                bounded_text(
                                    profile["header_line"],
                                    label="原始文件表头",
                                    formula_safe=False,
                                ),
                                int(profile["data_row_count"]),
                                int(profile["valid_row_count"]),
                                int(profile["parse_error_rows"]),
                                int(profile["blank_row_count"]),
                                sample_interval_s,
                                original_start_s,
                                original_end_s,
                                continued_start_s,
                                continued_end_s,
                                bounded_text(
                                    database_record.get("source_modified_utc"),
                                    label="来源修改时间",
                                ),
                                bounded_text(
                                    database_record.get("created_utc"),
                                    label="数据库收录时间",
                                ),
                            ),
                            (
                                "0",
                                None,
                                "0",
                                "0",
                                None,
                                None,
                                "0",
                                "0",
                                None,
                                "0",
                                None,
                                "0",
                                None,
                                None,
                                None,
                                "0",
                                "0",
                                "0",
                                "0",
                                "0.000000",
                                "0.000000",
                                "0.000000",
                                "0.000000",
                                "0.000000",
                                None,
                                None,
                            ),
                        )
                    )
                continued_offset_s += inclusive_duration_s
                del profile

        if raw_data_rows <= 0:
            raise StartStopWorkspaceError(
                "当前高亮材料没有可导出的完整原始数据。",
                422,
            )
        for raw_sheet_info in raw_sheet_infos:
            row_count = int(raw_sheet_info["rows"])
            raw_sheet_info["sheet"].auto_filter.ref = f"A1:L{row_count + 1}"
        for values, formats in raw_index_rows:
            append_formatted_row(
                raw_index_sheet,
                values,
                formats,
                force_text_columns=(13, 14),
            )
        raw_index_sheet.auto_filter.ref = f"A1:Z{len(raw_index_rows) + 1}"

        output = io.BytesIO()
        try:
            workbook.save(output)
            return output.getvalue()
        finally:
            if workbook_lifecycle is None:
                workbook.close()

    def _highlight_excel_temp_estimate(
        self,
        raw_context: dict[str, Any],
        processed_point_count: int,
    ) -> dict[str, int]:
        raw_rows = sum(
            max(0, _integer(source.get("snapshot_raw_row_count")))
            for series in raw_context.get("series", [])
            if isinstance(series, dict)
            for source in series.get("files", [])
            if isinstance(source, dict)
        )
        estimated = (
            self.EXCEL_TEMP_BASE_BYTES
            + raw_rows * self.EXCEL_TEMP_BYTES_PER_RAW_ROW
            + max(0, int(processed_point_count))
            * self.EXCEL_TEMP_BYTES_PER_PROCESSED_POINT
        )
        return {
            "raw_rows": raw_rows,
            "estimated_bytes": estimated,
            "required_free_bytes": estimated + self.EXCEL_TEMP_SAFETY_BYTES,
        }

    @contextlib.contextmanager
    def _highlight_excel_temp_scope(
        self,
        raw_context: dict[str, Any],
        processed_point_count: int,
    ) -> Iterable[list[Any]]:
        self.scratch_dir.mkdir(parents=True, exist_ok=True)
        if self.scratch_dir.is_symlink() or not self.scratch_dir.is_dir():
            raise StartStopWorkspaceError("Excel 导出临时目录无效。", 500)
        estimate = self._highlight_excel_temp_estimate(
            raw_context,
            processed_point_count,
        )
        try:
            free_bytes = int(shutil.disk_usage(self.scratch_dir).free)
        except OSError as exc:
            raise StartStopWorkspaceError(
                "无法检查 Excel 导出临时空间。",
                500,
            ) from exc
        if free_bytes < estimate["required_free_bytes"]:
            raise StartStopWorkspaceError(
                "Excel 导出临时空间不足，请减少高亮材料后再试。",
                507,
            )

        try:
            with self._scratch_budget.reserve(f"excel-{uuid.uuid4().hex}", estimate["required_free_bytes"]):
                export_temp = Path(
                    tempfile.mkdtemp(
                        prefix=".start-stop-excel-",
                        dir=str(self.scratch_dir),
                    )
                )
                previous_tempdir = tempfile.tempdir
                workbooks: list[Any] = []
                try:
                    # openpyxl write-only worksheets ask tempfile.gettempdir() for
                    # their XML streams. The chart-export lock makes this process-wide
                    # override exclusive, while all analysis jobs use explicit dirs.
                    tempfile.tempdir = str(export_temp)
                    yield workbooks
                finally:
                    for workbook in reversed(workbooks):
                        try:
                            workbook.close()
                        except Exception:
                            pass
                    tempfile.tempdir = previous_tempdir
                    shutil.rmtree(export_temp, ignore_errors=True)
        except ScratchSpaceError as exc:
            raise StartStopWorkspaceError(
                "Excel 导出临时空间不足；分析任务已预留空间，请稍后重试或减少高亮材料。", 507
            ) from exc

    @staticmethod
    def _chart_export_font_properties() -> tuple[Any, Any]:
        from matplotlib.font_manager import FontProperties

        regular_candidates = (
            os.environ.get("START_STOP_FONT_REGULAR", ""),
            "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
            "/Library/Fonts/NotoSansSC-VariableFont_wght.ttf",
            "/System/Library/Fonts/PingFang.ttc",
            "/System/Library/Fonts/STHeiti Light.ttc",
        )
        bold_candidates = (
            os.environ.get("START_STOP_FONT_BOLD", ""),
            "/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc",
            "/Library/Fonts/NotoSansSC-VariableFont_wght.ttf",
            "/System/Library/Fonts/PingFang.ttc",
            "/System/Library/Fonts/STHeiti Medium.ttc",
        )
        regular_path = next((path for path in regular_candidates if path and Path(path).is_file()), "")
        bold_path = next((path for path in bold_candidates if path and Path(path).is_file()), "")
        regular = FontProperties(fname=regular_path) if regular_path else FontProperties()
        bold = FontProperties(fname=bold_path, weight="bold") if bold_path else FontProperties(weight="bold")
        return regular, bold

    def _build_highlight_pdf(
        self,
        *,
        context: dict[str, Any],
        grouped: dict[tuple[str, str], list[dict[str, Any]]],
        names: dict[str, str],
        generated_at: dt.datetime,
        show_anomaly_markers: bool,
    ) -> bytes:
        try:
            from matplotlib.figure import Figure
        except ImportError as exc:
            raise StartStopWorkspaceError(
                "PDF 绘图组件当前不可用。",
                503,
            ) from exc

        regular_font, bold_font = self._chart_export_font_properties()
        figure = Figure(figsize=(11.69, 8.27), facecolor="white")
        axis = figure.add_subplot(111)
        plotted = 0
        for series_index, series_id in enumerate(context["requested"]):
            color = self.CHART_EXPORT_PALETTE[
                series_index % len(self.CHART_EXPORT_PALETTE)
            ]
            name = names.get(series_id) or str(
                context["known"][series_id].get("series_display_name") or series_id
            )
            for variant in context["variants"]:
                points = self._downsample(
                    grouped[(series_id, variant)],
                    self.PDF_EXPORT_POINTS_PER_CURVE,
                )
                if not points:
                    continue
                plotted += 1
                variant_label = self._chart_variant_label(variant)
                label = f"{name}（{variant_label}）" if context["mode"] == "compare" else name
                axis.plot(
                    [point["x"] for point in points],
                    [point["y"] for point in points],
                    color=color,
                    linewidth=2.0,
                    linestyle="--" if variant == "water" else "-",
                    label=label,
                )
                if (
                    show_anomaly_markers
                    and variant == "raw"
                    and context["metric"] != "overview"
                ):
                    normal = [point for point in points if point.get("status") == "normal"]
                    abnormal = [point for point in points if point.get("status") == "abnormal"]
                    if normal:
                        axis.scatter(
                            [point["x"] for point in normal],
                            [point["y"] for point in normal],
                            s=14,
                            color="#129B65",
                            edgecolors="white",
                            linewidths=0.35,
                            zorder=4,
                        )
                    if abnormal:
                        axis.scatter(
                            [point["x"] for point in abnormal],
                            [point["y"] for point in abnormal],
                            s=20,
                            color="#D92D20",
                            marker="x",
                            linewidths=0.9,
                            zorder=5,
                        )
        if plotted == 0:
            raise StartStopWorkspaceError("当前高亮曲线没有可导出的处理数据。")

        first_series = context["known"][context["requested"][0]]
        work_step_label = str(
            first_series.get("work_step_label") or context["work_step_key"]
        ).replace("⁻²", "$^{-2}$").replace("⁻", "-")
        figure.suptitle(
            f"{context.get('chart_title', '启停分析 · 当前高亮曲线')}\n{context['metric_spec']['label']} · {work_step_label}",
            fontproperties=bold_font,
            fontsize=15,
            y=0.97,
        )
        axis.set_xlabel(
            "累计时间 / h"
            if context["overview"] or context["x_axis"] == "time"
            else "循环数",
            fontproperties=regular_font,
            fontsize=11,
        )
        axis.set_ylabel(
            f"{context['metric_spec']['label']} / {context['metric_spec']['unit']}",
            fontproperties=regular_font,
            fontsize=11,
        )
        if context["metric"] == "minimum_time":
            axis.axhline(15.0, color="#D92D20", linewidth=1.2, linestyle=":")
        axis.grid(True, color="#DDE4E8", linewidth=0.65, alpha=0.85)
        axis.spines[["top", "right"]].set_visible(False)
        axis.tick_params(labelsize=9, colors="#344054")
        legend_columns = 1 if plotted == 1 else 2 if plotted <= 6 else 3 if plotted <= 15 else 4
        legend_rows = math.ceil(plotted / legend_columns)
        legend_font_size = 8.5 if plotted <= 6 else 7.5 if plotted <= 15 else 6.8
        legend = axis.legend(
            loc="upper center",
            bbox_to_anchor=(0.5, -0.16),
            ncol=legend_columns,
            frameon=False,
            prop=regular_font,
        )
        if legend is not None:
            for text in legend.get_texts():
                text.set_fontsize(legend_font_size)
        bottom_margin = min(0.42, 0.18 + legend_rows * 0.025)
        figure.subplots_adjust(
            left=0.09,
            right=0.98,
            top=0.84,
            bottom=bottom_margin,
        )
        footer = context.get("measurement_footer") or (
            "电位基准：Hg/HgO，不进行 RHE 换算；阴极末端：最后 1 s 中位数；"
            "异常：最低点 < 15 s。"
        )
        figure.text(
            0.01,
            0.02,
            footer,
            fontproperties=regular_font,
            fontsize=8,
            color="#475467",
        )
        output = io.BytesIO()
        figure.savefig(
            output,
            format="pdf",
            dpi=180,
            metadata={
                "Title": context.get("chart_title", "启停分析当前高亮曲线"),
                "Author": "Start-stop Studio",
                "Subject": context["metric_spec"]["label"],
                "CreationDate": generated_at,
            },
        )
        figure.clear()
        return output.getvalue()

    def chart_export(
        self,
        *,
        series_ids: Iterable[str],
        metric: str,
        x_axis: str,
        mode: str,
        export_format: str,
        show_anomaly_markers: bool = True,
    ) -> tuple[bytes, str, str]:
        export_format = str(export_format or "").strip().lower()
        if export_format not in self.CHART_EXPORT_FORMATS:
            raise StartStopWorkspaceError("导出格式只支持 xlsx 或 pdf。")
        if not self._chart_export_lock.acquire(blocking=False):
            raise StartStopWorkspaceError(
                "另一份高亮数据正在导出，请稍后重试。",
                409,
            )
        try:
            context = self._chart_request_context(
                series_ids=series_ids,
                metric=metric,
                x_axis=x_axis,
                mode=mode,
                maximum_series=self.MAX_HIGHLIGHTED_EXPORT_SERIES,
            )
            grouped, names, total_points = self._read_chart_groups(
                context,
                maximum_points=self.MAX_CHART_EXPORT_POINTS,
            )
            if total_points <= 0:
                raise StartStopWorkspaceError("当前高亮曲线没有可导出的处理数据。")
            stable = (
                self._file_cache_signature(context["path"])
                == context["data_signature"]
                and self._cache_generation_token() == context["generation"]
                and context["series_cache_key"][1] == context["generation"]
                and self._file_cache_signature(
                    self._path(self.SERIES_NAME, require=True)
                )
                == context["series_cache_key"][2]
            )
            if not stable:
                raise StartStopWorkspaceError(
                    "分析结果刚刚更新，请重新点击导出。",
                    409,
                )
            generated_at = dt.datetime.now(dt.timezone.utc)
            raw_context: dict[str, Any] | None = None
            if export_format == "xlsx":
                raw_context = self._raw_export_context(context)
                with self._highlight_excel_temp_scope(
                    raw_context,
                    total_points,
                ) as workbook_lifecycle:
                    data = self._build_highlight_excel(
                        context=context,
                        raw_context=raw_context,
                        grouped=grouped,
                        names=names,
                        generated_at=generated_at,
                        workbook_lifecycle=workbook_lifecycle,
                    )
                prefix = "启停分析_当前高亮_原始与处理数据"
                content_type = (
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                )
            else:
                data = self._build_highlight_pdf(
                    context=context,
                    grouped=grouped,
                    names=names,
                    generated_at=generated_at,
                    show_anomaly_markers=bool(show_anomaly_markers),
                )
                prefix = "启停分析_当前高亮_曲线图"
                content_type = "application/pdf"
            final_stable = (
                self._file_cache_signature(context["path"])
                == context["data_signature"]
                and self._cache_generation_token() == context["generation"]
                and context["series_cache_key"][1] == context["generation"]
                and self._file_cache_signature(
                    self._path(self.SERIES_NAME, require=True)
                )
                == context["series_cache_key"][2]
                and (
                    raw_context is None
                    or self._file_cache_signature(raw_context["snapshot_path"])
                    == raw_context["snapshot_signature"]
                )
            )
            if not final_stable:
                raise StartStopWorkspaceError(
                    "分析结果在导出期间已更新，请重新点击导出。",
                    409,
                )
            if len(data) > self.MAX_CHART_EXPORT_BYTES:
                raise StartStopWorkspaceError(
                    "导出文件过大，请减少高亮曲线后再试。",
                    413,
                )
            beijing = generated_at.astimezone(
                dt.timezone(dt.timedelta(hours=8))
            )
            filename = f"{prefix}_{beijing:%Y%m%d_%H%M%S}.{export_format}"
            return data, filename, content_type
        except StartStopWorkspaceError:
            raise
        except OSError as exc:
            if exc.errno == errno.ENOSPC:
                raise StartStopWorkspaceError(
                    "Excel 导出临时空间不足，请减少高亮材料后再试。",
                    507,
                ) from exc
            raise StartStopWorkspaceError(
                "无法生成当前高亮曲线导出文件。",
                500,
            ) from exc
        except Exception as exc:
            raise StartStopWorkspaceError(
                "无法生成当前高亮曲线导出文件。",
                500,
            ) from exc
        finally:
            self._chart_export_lock.release()

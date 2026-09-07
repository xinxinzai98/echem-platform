"""Highlighted LANBTS exports from immutable normalized records and full statistics."""
from __future__ import annotations

import csv
import datetime as dt
import errno
import io
import json
import math

from .start_stop_contracts import StartStopWorkspaceError
from .start_stop_stability import StabilityAnalysisError, MAX_SOURCE_BYTES, STATISTICS_VERSION

MAX_EXCEL_RECORDS = 2_000_000


def _identity(selected):
    return json.dumps([public for public, _ in selected], sort_keys=True, ensure_ascii=False)


def _number(value):
    try:
        value = float(value)
        return value if math.isfinite(value) else None
    except (TypeError, ValueError, OverflowError):
        return None


def _build_excel(workspace, selected, payload, lifecycle):
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Alignment, Font, PatternFill

    workbook = Workbook(write_only=True)
    lifecycle.append(workbook)
    workbook.properties.creator = "Start-stop Studio"
    workbook.properties.title = "蓝博当前高亮材料的完整记录与分析数据"
    fill = PatternFill("solid", fgColor="0B7F77")

    def append(sheet, values, *, literal=False, header=False):
        cells = []
        for value in values:
            if value is not None and (literal or isinstance(value, str)):
                value = str(value)
                if len(value) > 32767 or any(ord(c) < 32 and c not in "\t\n\r" for c in value):
                    raise StartStopWorkspaceError("原始记录含 Excel 无法无损保存的文本，请减少范围并检查源文件。", 422)
            cell = WriteOnlyCell(sheet, value=value)
            if isinstance(value, str):
                cell.data_type = "s"  # Preserve literal bytes-as-text; never evaluate source formulas.
            if header:
                cell.fill, cell.font = fill, Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cells.append(cell)
        sheet.append(cells)

    def table(name, headers):
        sheet = workbook.create_sheet(name)
        sheet.freeze_panes = "A2"
        for index in range(1, len(headers) + 1):
            from openpyxl.utils import get_column_letter
            sheet.column_dimensions[get_column_letter(index)].width = 24
        append(sheet, headers, header=True)
        return sheet

    notes = table("说明", ["项目", "说明"])
    for row in [
        ("数据边界", "仅导出当前高亮、同一工步的蓝博记录；不与 Hg/HgO 工作站数据混用纵轴。"),
        ("原始记录", "仪器 SDK 完整导出的标准化记录，保留全部字段及原始 µV/µA 数值列；以文本单元格无损保留，不抽样、不取端点、不做水位或参比换算。"),
        ("原始文件", "BTS 二进制仍保存在源数据库，通过父来源版本与 SHA-256 追溯；不将二进制嵌入工作表。"),
        ("处理数据", "当前指标的全量处理数据；全程指标逐条读取完整记录，循环指标来自全量循环计算，绝不从网页抽样点生成。"),
        ("接续规则", "沿用单份蓝博测试记录中的累计 time_s，不跨不同 BTS 测试自动接续。"),
        ("统计版本", STATISTICS_VERSION),
    ]:
        append(notes, row)
    index_sheet = table("材料与来源", ["编号", "完整材料名", "测试模式", "工步", "源文件", "标准化 SHA-256", "父来源 SHA-256", "标准化版本", "父来源版本", "完整记录数"])
    summary_sheet = table("统计汇总", ["编号", "完整材料名", "统计项", "数值"])
    total_rows = 0
    metric = payload["metric"]
    for index, (public, raw) in enumerate(selected, 1):
        prepared = workspace.stability_analyzer._prepared_source(public, raw)
        content = workspace.source_database.read_source_export_content(
            source_version_id=public["source_version_id"], expected_sha256=public["sha256"],
            expected_size_bytes=public["size_bytes"], maximum_file_bytes=MAX_SOURCE_BYTES,
        )
        reader = csv.reader(io.StringIO(content.decode("utf-8"), newline=""))
        headers = next(reader)
        raw_sheet = table(f"{index:02d}_原始记录_1", headers)
        curve_sheet = table(f"{index:02d}_处理数据_1", [payload["metric_spec"]["x_label"], f'{payload["metric_spec"]["label"]} / {payload["metric_spec"]["unit"]}', "正常或异常", "累计时间 / h"])
        raw_count, curve_count = 0, 0
        raw_chunk, curve_chunk = 1, 1
        raw_sheet_rows = curve_sheet_rows = 0

        def write_curve(x, y, status, time_h):
            nonlocal curve_sheet, curve_count, curve_chunk, curve_sheet_rows
            if x is None or y is None:
                return
            if curve_sheet_rows >= workspace.MAX_EXCEL_DATA_ROWS_PER_SHEET:
                curve_chunk += 1
                curve_sheet = table(f"{index:02d}_处理数据_{curve_chunk}", [payload["metric_spec"]["x_label"], f'{payload["metric_spec"]["label"]} / {payload["metric_spec"]["unit"]}', "正常或异常", "累计时间 / h"])
                curve_sheet_rows = 0
            append(curve_sheet, [x, y, status, time_h])
            curve_count += 1
            curve_sheet_rows += 1

        for values in reader:
            total_rows += 1
            if total_rows > MAX_EXCEL_RECORDS:
                raise StartStopWorkspaceError("完整记录过多，请分批选择高亮材料导出。", 413)
            if raw_sheet_rows >= workspace.MAX_EXCEL_DATA_ROWS_PER_SHEET:
                raw_chunk += 1
                raw_sheet = table(f"{index:02d}_原始记录_{raw_chunk}", headers)
                raw_sheet_rows = 0
            append(raw_sheet, values, literal=True)
            raw_count += 1
            raw_sheet_rows += 1
            if metric in {"overview", "current", "current_density"}:
                row = dict(zip(headers, values))
                time_s = _number(row.get("time_s"))
                # Match the analyzer's validity gate without altering the raw sheet.
                if any(_number(row.get(key)) is None for key in ("time_s", "potential_v", "current_ma")):
                    continue
                field = {"overview":"potential_v", "current":"current_ma", "current_density":"current_density_ma_cm2"}[metric]
                write_curve(time_s/3600, _number(row.get(field)), "", time_s/3600)
        if metric not in {"overview", "current", "current_density"}:
            field = {"stress_endpoint":"stress_endpoint_v", "recovery_endpoint":"recovery_endpoint_v", "negative_shift":"negative_shift_mv", "minimum_time":"minimum_time_s"}[metric]
            for cycle in prepared["cycles"]:
                write_curve(cycle["cycle"], cycle[field], cycle["status"], cycle["time_h"])
        for key, value in prepared["summary"].items():
            append(summary_sheet, [index, public["display_name"], key, value])
        append(index_sheet, [index, public["display_name"], public["analysis_mode_label"], public["protocol_label"],
            public["source_file"], public["sha256"], public["parent_sha256"], public["source_version_id"], public["parent_source_version_id"], raw_count])
        if curve_count == 0:
            append(notes, [public["display_name"], "当前指标没有有效处理点；完整原始记录仍已保留。"])
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def export_stability(workspace, *, series_ids, analysis_mode, metric, export_format):
    if export_format not in {"xlsx", "pdf"}:
        raise StartStopWorkspaceError("导出格式只支持 xlsx 或 pdf。")
    requested = list(series_ids)
    if not workspace._chart_export_lock.acquire(blocking=False):
        raise StartStopWorkspaceError("另一份高亮数据正在导出，请稍后重试。", 409)
    try:
        analyzer = workspace.stability_analyzer
        known = {public["series_id"]:(public, raw) for public, raw in analyzer._sources()}
        if any(key not in known for key in requested):
            raise StartStopWorkspaceError("蓝博记录不存在或已更新。", 404)
        selected = [known[key] for key in requested]
        identity = _identity(selected)
        for public, raw in selected:
            metadata = raw["metadata"]
            parent = workspace.source_database.source_version_identity(public["parent_source_version_id"])
            if not parent or parent["sha256"] != public["parent_sha256"] or parent["repository_path"] != metadata.get("parent_repository_path"):
                raise StartStopWorkspaceError("蓝博原始 BTS 来源无法核验，请先重新下载该测试记录。", 409)
            if metadata.get("parent_size_bytes") is not None and parent["size_bytes"] != metadata["parent_size_bytes"]:
                raise StartStopWorkspaceError("蓝博父来源大小与记录不一致，请重新下载。", 409)
        # Reuse the same protocol and scientific gates as the plot endpoint.
        payload = analyzer.chart(series_ids=requested, analysis_mode=analysis_mode, metric=metric, max_points=12_000)
        after_chart = {public["series_id"]:(public, raw) for public, raw in analyzer._sources()}
        if any(key not in after_chart for key in requested) or _identity([after_chart[key] for key in requested]) != identity:
            raise StartStopWorkspaceError("蓝博记录刚刚更新，请重新导出。", 409)
        if any(public["downsample_stride"] > 1 for public, _ in selected):
            raise StartStopWorkspaceError("旧版抽样记录不能导出为完整分析，请先重新下载全量记录。", 422)
        # The preview and selected normalized versions must describe identical bytes.
        for item, (public, _) in zip(payload["series"], selected):
            if item["provenance"]["normalized_sha256"] != public["sha256"]:
                raise StartStopWorkspaceError("蓝博记录刚刚更新，请重新导出。", 409)
        generated_at = dt.datetime.now(dt.timezone.utc)
        if export_format == "xlsx":
            estimated_rows = sum(max(public["record_count"], public["point_count"], public["size_bytes"] // 24) for public, _ in selected)
            if estimated_rows > MAX_EXCEL_RECORDS:
                raise StartStopWorkspaceError("完整记录过多，请分批选择高亮材料导出。", 413)
            raw_context = {"series":[{"files":[{"snapshot_raw_row_count":estimated_rows}]}]}
            with workspace._highlight_excel_temp_scope(raw_context, estimated_rows) as lifecycle:
                data = _build_excel(workspace, selected, payload, lifecycle)
            mime = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        else:
            context = {"requested":requested, "known":{key:{"work_step_label":known[key][0]["protocol_label"]} for key in requested},
                "work_step_key":selected[0][0]["protocol_key"], "variants":["raw"], "mode":"raw", "metric":metric,
                "metric_spec":payload["metric_spec"], "overview":payload["metric_spec"]["x_key"] == "time_h",
                "x_axis":"time" if payload["metric_spec"]["x_key"] == "time_h" else "cycle",
                "chart_title":f'蓝博{payload["analysis_mode_label"]} · 当前高亮曲线',
                "measurement_footer":"电压按 LANBTS 原始标尺；参照未确认，不进行 RHE 或水位补偿；PDF 为抽样绘图，统计由完整记录计算。"}
            data = workspace._build_highlight_pdf(context=context,
                grouped={(item["series_id"],"raw"):item["points"] for item in payload["series"]},
                names={item["series_id"]:item["display_name"] for item in payload["series"]},
                generated_at=generated_at, show_anomaly_markers=True)
            mime = "application/pdf"
        current = {public["series_id"]:(public, raw) for public, raw in analyzer._sources()}
        if any(key not in current for key in requested) or _identity([current[key] for key in requested]) != identity:
            raise StartStopWorkspaceError("蓝博记录在导出期间已更新，请重新导出。", 409)
        if len(data) > workspace.MAX_CHART_EXPORT_BYTES:
            raise StartStopWorkspaceError("导出文件过大，请减少高亮材料后重试。", 413)
        local_time = generated_at.astimezone(dt.timezone(dt.timedelta(hours=8)))
        return data, f"蓝博_当前高亮_原始与处理数据_{local_time:%Y%m%d_%H%M%S}.{export_format}", mime
    except StabilityAnalysisError as exc:
        raise StartStopWorkspaceError(str(exc), exc.status) from exc
    except ImportError as exc:
        raise StartStopWorkspaceError("图表导出组件未安装。", 503) from exc
    except OSError as exc:
        raise StartStopWorkspaceError("导出临时空间不足。" if exc.errno == errno.ENOSPC else "导出文件读取失败。", 507 if exc.errno == errno.ENOSPC else 503) from exc
    finally:
        workspace._chart_export_lock.release()

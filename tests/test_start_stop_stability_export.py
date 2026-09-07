import copy
import io
import tempfile
import unittest
from pathlib import Path
from unittest import mock

from openpyxl import load_workbook

from echem_platform.start_stop import StartStopWorkspace, StartStopWorkspaceError
from echem_platform.start_stop_database import StartStopDatabase
from tests.test_start_stop_stability import FakeDatabase


class StabilityExportTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary.name)
        self.database = StartStopDatabase(self.root / "repository.sqlite3")
        fake = FakeDatabase()
        for index, source in enumerate(fake.rows):
            raw = self.root / f"raw-{index}.bts"
            raw.write_bytes(f"synthetic raw fixture {index}".encode())
            parent = self.database.ingest_staged_file(raw, {
                "machine_id":"fixture", "root_label":"data", "remote_path":raw.name,
                "repository_path":raw.name, "size":raw.stat().st_size, "last_write_ticks":index+1,
                "is_candidate":False,
            })
            content = fake.contents[source["source_version_id"]]
            path = self.root / f"normalized-{index}.csv"
            path.write_bytes(content)
            self.database.ingest_staged_file(path, {
                "machine_id":"fixture", "root_label":"data", "remote_path":str(path.name),
                "repository_path":str(path.name), "size":len(content), "last_write_ticks":index+1,
                "is_candidate":True, "candidate_kind":source["candidate_kind"], **source["metadata"],
                "record_count":80 if index == 0 else 61, "exported_point_count":80 if index == 0 else 61,
                "parent_source_version_id":parent["version_id"], "parent_repository_path":raw.name,
                "parent_sha256":parent["sha256"], "parent_size_bytes":parent["size_bytes"],
            })
        self.workspace = StartStopWorkspace(self.database, self.root / "published", repository_mode=True, scratch_dir=self.root / "scratch")
        self.ids = {row["analysis_mode"]:row["series_id"] for row in self.workspace.stability_catalog()["series"]}

    def tearDown(self):
        self.temporary.cleanup()

    def export(self, metric="stress_endpoint", export_format="xlsx", mode="start_stop"):
        return self.workspace.stability_chart_export(series_ids=[self.ids[mode]], analysis_mode=mode, metric=metric, export_format=export_format)

    def test_full_raw_records_and_numeric_cycle_results_preserve_source(self):
        before = (self.root / "normalized-0.csv").read_bytes()
        data, filename, mime = self.export()
        self.assertTrue(filename.endswith(".xlsx"))
        self.assertIn("spreadsheetml", mime)
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
        try:
            raw = list(workbook["01_原始记录_1"].values)
            processed = list(workbook["01_处理数据_1"].values)
            self.assertEqual(len(raw), 81)
            self.assertEqual(raw[0][:3], ("time_s", "potential_v", "current_ma"))
            self.assertEqual(raw[1][:3], ("0", "-2.0", "-300.0"))
            self.assertEqual(workbook["01_原始记录_1"]["B2"].data_type, "s")
            self.assertEqual(len(processed), 3)
            self.assertAlmostEqual(processed[1][1], -1.5185)
            self.assertEqual(processed[1][2], "abnormal")
            self.assertIn("BTS 二进制", str(list(workbook["说明"].values)))
        finally:
            workbook.close()
        self.assertEqual((self.root / "normalized-0.csv").read_bytes(), before)
        self.assertEqual(list((self.root / "scratch").glob(".start-stop-excel-*")), [])

    def test_overview_export_uses_all_rows_not_the_cached_preview(self):
        with mock.patch("echem_platform.start_stop_stability.MAX_POINTS", 10):
            preview = self.workspace.stability_chart(series_ids=[self.ids["start_stop"]], analysis_mode="start_stop", metric="overview")
            self.assertEqual(len(preview["series"][0]["points"]), 10)
            data, _, _ = self.export("overview")
        workbook = load_workbook(io.BytesIO(data), read_only=True)
        try:
            self.assertEqual(len(list(workbook["01_处理数据_1"].values)), 81)
        finally:
            workbook.close()

    def test_pdf_uses_shared_renderer_with_lanbts_measurement_boundary(self):
        with mock.patch.object(self.workspace, "_build_highlight_pdf", wraps=self.workspace._build_highlight_pdf) as render:
            data, filename, mime = self.export(export_format="pdf")
        self.assertEqual(mime, "application/pdf")
        self.assertTrue(data.startswith(b"%PDF"))
        self.assertTrue(filename.endswith(".pdf"))
        context = render.call_args.kwargs["context"]
        self.assertIn("LANBTS 原始标尺", context["measurement_footer"])
        self.assertNotIn("电位基准：Hg/HgO", context["measurement_footer"])
        self.assertEqual(context["variants"], ["raw"])

    def test_rejects_old_sampling_and_concurrent_source_change(self):
        sources = self.workspace.stability_analyzer._sources()
        legacy = copy.deepcopy(sources)
        legacy[0][0]["downsample_stride"] = 2
        # Identify the intended start-stop fixture independently of catalog order.
        for public, raw in legacy:
            if public["analysis_mode"] == "start_stop":
                public["downsample_stride"] = 2
        with mock.patch.object(self.workspace.stability_analyzer, "_sources", return_value=legacy):
            with self.assertRaises(StartStopWorkspaceError) as caught:
                self.export("overview")
        self.assertEqual(caught.exception.status, 422)

        changed = False
        def current_sources():
            value = copy.deepcopy(sources)
            if changed:
                for public, _ in value:
                    public["display_name"] += " updated"
            return value
        def build(*args):
            nonlocal changed
            changed = True
            return b"not released"
        with mock.patch.object(self.workspace.stability_analyzer, "_sources", side_effect=current_sources), mock.patch("echem_platform.start_stop_stability_export._build_excel", side_effect=build):
            with self.assertRaises(StartStopWorkspaceError) as caught:
                self.export()
        self.assertEqual(caught.exception.status, 409)
        self.assertFalse(self.workspace._chart_export_lock.locked())

    def test_rejects_unverifiable_parent_before_export(self):
        with mock.patch.object(self.database, "source_version_identity", return_value=None):
            with self.assertRaises(StartStopWorkspaceError) as caught:
                self.export()
        self.assertEqual(caught.exception.status, 409)
